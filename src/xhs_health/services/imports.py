import csv
import math
from datetime import date
from io import BytesIO, StringIO

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from xhs_health.models import Account, AccountDailySnapshot, ImportBatch, ImportErrorRow, Note, NoteDailyMetric
from xhs_health.schemas import AccountImportIn, ImportAccountsResponse, validate_import_payload


def _jsonable(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, list):
        return [_jsonable(item) for item in value]
    if isinstance(value, dict):
        return {key: _jsonable(item) for key, item in value.items()}
    return value


def _interaction_rate(payload: dict) -> float | None:
    reads = payload.get("read_count")
    if not reads:
        return None
    interactions = sum(
        payload.get(key) or 0 for key in ("like_count", "collect_count", "comment_count", "share_count")
    )
    return interactions / reads


def _cqi(payload: dict) -> float | None:
    reads = payload.get("read_count")
    if not reads:
        return None
    weighted = (
        0.25 * (payload.get("like_count") or 0)
        + 0.35 * (payload.get("collect_count") or 0)
        + 0.25 * (payload.get("comment_count") or 0)
        + 0.15 * (payload.get("share_count") or 0)
    )
    return weighted / reads


def import_accounts(
    session: Session,
    accounts: list[AccountImportIn],
    batch: ImportBatch | None = None,
    row_numbers: list[int] | None = None,
) -> ImportAccountsResponse:
    """Import JSON-style accounts.

    When `batch` is provided, validation errors are recorded as ImportErrorRow
    entries and the batch is finalised with status / error counts. This unifies
    JSON import with the CSV/XLSX path so both share the same audit trail.
    `row_numbers` maps each payload to its original source row (e.g. the first
    CSV row of an aggregated account); without it, payloads are numbered from 1.
    """
    counts = {
        "accounts_upserted": 0,
        "snapshots_upserted": 0,
        "notes_upserted": 0,
        "note_metrics_upserted": 0,
    }
    error_rows = 0

    for index, payload in enumerate(accounts):
        row_number = (
            row_numbers[index] if row_numbers and index < len(row_numbers) else index + 1
        )
        errors = validate_import_payload(payload)
        errors.extend(_note_ownership_errors(session, payload))
        if errors:
            error_rows += 1
            if batch is not None:
                for field_name, message in errors:
                    session.add(
                        ImportErrorRow(
                            batch_id=batch.id,
                            row_number=row_number,
                            field_name=field_name,
                            message=message,
                            raw_payload=_jsonable(payload.model_dump()),
                        )
                    )
            continue

        try:
            account = _upsert_account(session, payload)
            counts["accounts_upserted"] += 1

            for snapshot_payload in payload.snapshots:
                snapshot_data = snapshot_payload.model_dump()
                _upsert_snapshot(session, account.id, snapshot_data)
                counts["snapshots_upserted"] += 1

            for note_payload in payload.notes:
                note = _upsert_note(session, account.id, note_payload.model_dump(exclude={"metrics"}))
                counts["notes_upserted"] += 1
                for metric_payload in note_payload.metrics:
                    metric_data = metric_payload.model_dump()
                    _upsert_note_metric(session, account.id, note.id, note.note_id, metric_data)
                    counts["note_metrics_upserted"] += 1
        except ValueError as exc:
            # Business-rule violation inside upsert (e.g. note_id owned by another
            # account): record against the offending row instead of failing silent.
            error_rows += 1
            if batch is not None:
                session.add(
                    ImportErrorRow(
                        batch_id=batch.id,
                        row_number=row_number,
                        field_name="note_id",
                        message=str(exc),
                        raw_payload=_jsonable(payload.model_dump()),
                    )
                )
        except Exception as exc:  # pragma: no cover - defensive
            error_rows += 1
            if batch is not None:
                session.add(
                    ImportErrorRow(
                        batch_id=batch.id,
                        row_number=row_number,
                        field_name="__row__",
                        message=f"unexpected error: {exc}",
                        raw_payload=_jsonable(payload.model_dump()),
                    )
                )

    response = ImportAccountsResponse(**counts)
    if batch is not None:
        batch.valid_rows = counts["accounts_upserted"]
        batch.error_rows = error_rows
        batch.status = "completed_with_errors" if error_rows else "completed"
        response.import_batch_id = batch.id
        response.total_rows = batch.total_rows
        response.error_rows = error_rows
    return response


def import_flat_file(session: Session, filename: str, content: bytes) -> ImportAccountsResponse:
    try:
        rows = _read_rows(filename, content)
    except HTTPException:
        raise
    except Exception as exc:  # garbage .xlsx bytes, non-UTF-8 text, corrupt workbook
        raise HTTPException(
            status_code=422,
            detail=f"import file could not be parsed as UTF-8 CSV or XLSX: {type(exc).__name__}",
        ) from exc
    batch = ImportBatch(filename=filename, total_rows=len(rows), valid_rows=0, error_rows=0)
    session.add(batch)
    session.flush()

    accounts: dict[str, dict] = {}
    account_first_row: dict[str, int] = {}
    row_level_errors = 0
    for index, row in enumerate(rows, start=2):
        errors = _validate_row(row)
        if errors:
            for field_name, message in errors:
                session.add(
                    ImportErrorRow(
                        batch_id=batch.id,
                        row_number=index,
                        field_name=field_name,
                        message=message,
                        raw_payload=_jsonable(row),
                    )
                )
            row_level_errors += 1
            continue

        platform_uid = str(row.get("platform_uid") or "").strip()
        if not platform_uid:
            continue
        account = accounts.setdefault(
            platform_uid,
            {
                "platform_uid": platform_uid,
                "nickname": str(row.get("nickname") or platform_uid),
                "category": _empty_to_none(row.get("category")),
                "fan_quality_score": _float_or_none(row.get("fan_quality_score")),
                "cpe": _float_or_none(row.get("cpe")),
                "avg_cpe_benchmark": _float_or_none(row.get("avg_cpe_benchmark")),
                "business_stability": _float_or_none(row.get("business_stability")),
                "violation_count_180d": _int_or_none(row.get("violation_count_180d")),
                "ad_compliance_rate": _float_or_none(row.get("ad_compliance_rate")),
                "audit_pass_rate": _float_or_none(row.get("audit_pass_rate")),
                "shadowban_risk": _float_or_none(row.get("shadowban_risk")),
                "snapshots": [],
                "notes": [],
            },
        )
        account_first_row.setdefault(platform_uid, index)
        if row.get("data_date"):
            account["snapshots"].append(
                {
                    "data_date": row["data_date"],
                    "fans_count": _int_or_none(row.get("fans_count")),
                    "fans_delta": _int_or_none(row.get("fans_delta")),
                    "notes_count": _int_or_none(row.get("notes_count")),
                    "total_reads": _int_or_none(row.get("total_reads")),
                    "total_likes": _int_or_none(row.get("total_likes")),
                    "total_collects": _int_or_none(row.get("total_collects")),
                    "total_comments": _int_or_none(row.get("total_comments")),
                    "total_shares": _int_or_none(row.get("total_shares")),
                    "publish_count": _int_or_none(row.get("publish_count")),
                    "data_source": str(row.get("data_source") or "file"),
                }
            )
        note_id = _empty_to_none(row.get("note_id"))
        if note_id:
            account["notes"].append(
                {
                    "note_id": note_id,
                    "title": _empty_to_none(row.get("note_title")),
                    "content_type": str(row.get("content_type") or "image"),
                    "is_ad": _bool_or_false(row.get("is_ad")),
                    "is_original": not _bool_or_false(row.get("is_repost")),
                    "tags": _split_tags(row.get("tags")),
                    "metrics": [
                        {
                            "data_date": row["data_date"],
                            "read_count": _int_or_none(row.get("read_count")),
                            "like_count": _int_or_none(row.get("like_count")),
                            "collect_count": _int_or_none(row.get("collect_count")),
                            "comment_count": _int_or_none(row.get("comment_count")),
                            "share_count": _int_or_none(row.get("share_count")),
                            "data_source": str(row.get("data_source") or "file"),
                        }
                    ],
                }
            )

    # Run business-rule validation per aggregated account; record failures with the
    # original CSV row number of that account's first appearance. This catches
    # violations (e.g. ratio out of [0,1]) that _validate_row's type checks miss.
    valid_payloads: list[AccountImportIn] = []
    valid_row_numbers: list[int] = []
    account_level_errors = 0
    for platform_uid, raw in accounts.items():
        try:
            payload = AccountImportIn.model_validate(raw)
        except Exception as exc:
            account_level_errors += 1
            session.add(
                ImportErrorRow(
                    batch_id=batch.id,
                    row_number=account_first_row[platform_uid],
                    field_name="__row__",
                    message=f"account payload rejected: {exc}",
                    raw_payload=_jsonable(raw),
                )
            )
            continue
        rule_errors = validate_import_payload(payload)
        if rule_errors:
            account_level_errors += 1
            for field_name, message in rule_errors:
                session.add(
                    ImportErrorRow(
                        batch_id=batch.id,
                        row_number=account_first_row[platform_uid],
                        field_name=field_name,
                        message=message,
                        raw_payload=_jsonable(raw),
                    )
                )
            continue
        valid_payloads.append(payload)
        valid_row_numbers.append(account_first_row[platform_uid])

    # Pass the batch through so upsert-time failures (e.g. a note_id owned by
    # another account) also land in the batch error trail.
    result = import_accounts(session, valid_payloads, batch=batch, row_numbers=valid_row_numbers)
    batch.valid_rows = result.accounts_upserted
    batch.error_rows = row_level_errors + account_level_errors + result.error_rows
    batch.status = "completed_with_errors" if batch.error_rows else "completed"
    result.import_batch_id = batch.id
    result.total_rows = batch.total_rows
    result.error_rows = batch.error_rows
    return result


def _finite_number_or_none(value):
    """Parse a user-supplied number; reject inf/nan (e.g. "1e999", "nan").

    Returns (parsed, error_message). int(float("1e999")) raises OverflowError
    and float("nan") poisons scores silently, so both are caught here.
    """
    value = _empty_to_none(value)
    if value is None:
        return None, None
    try:
        parsed = float(value)
    except (TypeError, ValueError, OverflowError):
        return None, "must be a number"
    if not math.isfinite(parsed):
        return None, "must be a finite number"
    return parsed, None


def _validate_row(row: dict) -> list[tuple[str, str]]:
    errors: list[tuple[str, str]] = []
    if not _empty_to_none(row.get("platform_uid")):
        errors.append(("platform_uid", "platform_uid is required"))

    if _empty_to_none(row.get("data_date")):
        try:
            date.fromisoformat(str(row["data_date"]))
        except ValueError:
            errors.append(("data_date", "data_date must be YYYY-MM-DD"))

    for field_name in (
        "fans_count",
        "fans_delta",
        "notes_count",
        "total_reads",
        "total_likes",
        "total_collects",
        "total_comments",
        "total_shares",
        "publish_count",
        "violation_count_180d",
        "read_count",
        "like_count",
        "collect_count",
        "comment_count",
        "share_count",
    ):
        if _empty_to_none(row.get(field_name)) is not None:
            _, message = _finite_number_or_none(row[field_name])
            if message is None:
                try:
                    int(float(row[field_name]))
                except (TypeError, ValueError, OverflowError):
                    message = "must be an integer"
            if message is not None:
                errors.append((field_name, f"{field_name} {message}"))

    for field_name in (
        "ad_compliance_rate",
        "audit_pass_rate",
        "shadowban_risk",
        "fan_quality_score",
        "cpe",
        "avg_cpe_benchmark",
        "business_stability",
    ):
        if _empty_to_none(row.get(field_name)) is not None:
            _, message = _finite_number_or_none(row[field_name])
            if message is not None:
                errors.append((field_name, f"{field_name} {message}"))

    if _empty_to_none(row.get("note_id")) and not _empty_to_none(row.get("data_date")):
        errors.append(("data_date", "note metric rows require data_date"))
    return errors


def _read_rows(filename: str, content: bytes) -> list[dict]:
    lowered = filename.lower()
    if lowered.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        return [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
            for row in rows[1:]
        ]
    text = content.decode("utf-8-sig")
    return list(csv.DictReader(StringIO(text)))


def _empty_to_none(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _float_or_none(value):
    value = _empty_to_none(value)
    return None if value is None else float(value)


def _int_or_none(value):
    value = _empty_to_none(value)
    return None if value is None else int(float(value))


def _bool_or_false(value):
    value = _empty_to_none(value)
    if value is None:
        return False
    return str(value).strip().lower() in {"1", "true", "yes", "y", "是"}


def _split_tags(value):
    value = _empty_to_none(value)
    if value is None:
        return []
    return [item.strip() for item in str(value).replace("，", ",").split(",") if item.strip()]


def _note_ownership_errors(session: Session, payload: AccountImportIn) -> list[tuple[str, str]]:
    """Reject note_ids already owned by a different account BEFORE any writes.

    Checking up-front keeps the payload atomic: rejecting after the account and
    snapshots were written would leave a half-imported row behind.
    """
    note_ids = [note.note_id for note in payload.notes]
    if not note_ids:
        return []
    target = session.scalar(select(Account.id).where(Account.platform_uid == payload.platform_uid))
    owners = {
        note_id: account_id
        for note_id, account_id in session.execute(
            select(Note.note_id, Note.account_id).where(Note.note_id.in_(note_ids))
        ).all()
    }
    errors: list[tuple[str, str]] = []
    for note_id, owner_id in owners.items():
        if target is None or owner_id != target:
            errors.append(
                (
                    "note_id",
                    f"note_id {note_id} already belongs to account {owner_id}",
                )
            )
    return errors


def _merge_non_none(target, data: dict, raw_payload: dict | None = None) -> None:
    """Update `target` with `data`, skipping None values.

    A later import that omits a column must not wipe what an earlier import
    provided (incremental daily imports routinely lack compliance columns).
    """
    for key, value in data.items():
        if value is not None:
            setattr(target, key, value)
    if raw_payload is not None:
        merged = dict(getattr(target, "raw_payload", None) or {})
        merged.update({key: value for key, value in data.items() if value is not None})
        target.raw_payload = _jsonable(merged)


def _upsert_account(session: Session, payload: AccountImportIn) -> Account:
    account = session.scalar(select(Account).where(Account.platform_uid == payload.platform_uid))
    data = payload.model_dump(exclude={"snapshots", "notes"})
    if account is None:
        account = Account(**data)
        session.add(account)
        session.flush()
        return account
    _merge_non_none(account, data)
    session.flush()
    return account


def _upsert_snapshot(session: Session, account_id: int, data: dict) -> AccountDailySnapshot:
    snapshot = session.scalar(
        select(AccountDailySnapshot).where(
            AccountDailySnapshot.account_id == account_id,
            AccountDailySnapshot.data_date == data["data_date"],
            AccountDailySnapshot.data_source == data["data_source"],
        )
    )
    if snapshot is None:
        snapshot = AccountDailySnapshot(account_id=account_id, raw_payload=_jsonable(data), **data)
        session.add(snapshot)
        session.flush()
        return snapshot
    _merge_non_none(snapshot, data, raw_payload=_jsonable(data))
    session.flush()
    return snapshot


def _upsert_note(session: Session, account_id: int, data: dict) -> Note:
    note = session.scalar(select(Note).where(Note.note_id == data["note_id"]))
    if note is not None and note.account_id != account_id:
        # note_id is globally unique: never silently reassign a note to another account.
        raise ValueError(
            f"note_id {data['note_id']} already belongs to account {note.account_id}"
        )
    if note is None:
        note = Note(account_id=account_id, **data)
        session.add(note)
        session.flush()
        return note
    _merge_non_none(note, data)
    session.flush()
    return note


def _upsert_note_metric(
    session: Session,
    account_id: int,
    note_pk: int,
    note_id: str,
    data: dict,
) -> NoteDailyMetric:
    metric = session.scalar(
        select(NoteDailyMetric).where(
            NoteDailyMetric.note_id == note_id,
            NoteDailyMetric.data_date == data["data_date"],
            NoteDailyMetric.data_source == data["data_source"],
        )
    )
    data["interaction_rate"] = _interaction_rate(data)
    data["cqi"] = _cqi(data)
    if metric is None:
        metric = NoteDailyMetric(
            account_id=account_id,
            note_pk=note_pk,
            note_id=note_id,
            raw_payload=_jsonable(data),
            **data,
        )
        session.add(metric)
        session.flush()
        return metric
    _merge_non_none(metric, data, raw_payload=_jsonable(data))
    session.flush()
    return metric
